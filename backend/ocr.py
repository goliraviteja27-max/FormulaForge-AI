from __future__ import annotations

import io
from pathlib import Path

import fitz
from docx import Document
from PIL import Image
from pypdf import PdfReader

_reader = None


def _easyocr_reader():
    global _reader
    if _reader is None:
        import easyocr

        _reader = easyocr.Reader(["en"], gpu=False)
    return _reader


def extract_image_text(image_bytes: bytes) -> str:
    image = Image.open(io.BytesIO(image_bytes)).convert("RGB")
    return "\n".join(_easyocr_reader().readtext(image, detail=0, paragraph=True))


def extract_pdf_text(file_bytes: bytes, max_pages: int = 35) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    typed = "\n".join((page.extract_text() or "") for page in reader.pages[:max_pages]).strip()
    # A small text layer usually means the PDF is a scan, so read rendered pages.
    if len(typed) >= 120:
        return typed

    document = fitz.open(stream=file_bytes, filetype="pdf")
    lines: list[str] = []
    for page_index in range(min(max_pages, document.page_count)):
        page = document.load_page(page_index)
        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        lines.append(extract_image_text(pix.tobytes("png")))
    return "\n".join(lines)


def extract_docx_text(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(part for part in parts if part.strip())


def extract_text(filename: str, file_bytes: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        return extract_pdf_text(file_bytes)
    if extension == ".docx":
        return extract_docx_text(file_bytes)
    if extension == ".txt":
        return file_bytes.decode("utf-8", errors="replace")
    if extension in {".png", ".jpg", ".jpeg"}:
        return extract_image_text(file_bytes)
    raise ValueError(f"Unsupported file type: {extension or 'unknown'}")
